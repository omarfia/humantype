"""Tests for the markdown -> docx formatting renderer."""

from docx import Document

from humantype.extract import extract_paragraphs
from humantype.formatting import get_formatting_docx, markdown_text_to_docx


def test_markdown_renders_headings_bold_italic_lists(tmp_path):
    md = (
        "# Title\n\n"
        "Some **bold** and *italic* text with `code`.\n\n"
        "- item one\n"
        "- item two\n\n"
        "> a quote\n"
    )
    out = tmp_path / "md.docx"
    markdown_text_to_docx(md, str(out))

    document = Document(str(out))
    paragraphs = document.paragraphs
    assert paragraphs[0].style.name == "Title" or "Heading" in paragraphs[0].style.name
    assert paragraphs[0].text == "Title"

    body = paragraphs[1]
    assert body.text == "Some bold and italic text with code."
    assert any(r.bold for r in body.runs)
    assert any(r.italic for r in body.runs)

    bullets = [p for p in paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullets] == ["item one", "item two"]

    quotes = [p for p in paragraphs if p.style.name == "Quote"]
    assert quotes[0].text == "a quote"


def test_markdown_docx_text_matches_plain_extraction(tmp_path):
    md_path = tmp_path / "sample.md"
    md_path.write_text(
        "# Title\n\nThe **quick** fox and [a link](http://x).\n\n- one\n- two\n",
        encoding="utf-8",
    )
    docx_path = get_formatting_docx(str(md_path))
    rendered = extract_paragraphs(docx_path)
    assert rendered[0] == "Title"
    assert rendered[1] == "The quick fox and a link."
    assert rendered[2] == "one"
    assert rendered[3] == "two"


def test_docx_source_is_used_directly(tmp_path):
    docx_path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Hello world")
    document.save(str(docx_path))
    assert get_formatting_docx(str(docx_path)) == str(docx_path)


def test_docx_to_html_preserves_formatting(tmp_path):
    from humantype.gdocs_format import docx_to_html

    docx_path = tmp_path / "styled.docx"
    document = Document()
    document.add_heading("Title", level=1)
    para = document.add_paragraph()
    run = para.add_run("bold")
    run.bold = True
    para.add_run(" and ")
    run = para.add_run("italic")
    run.italic = True
    document.add_paragraph("item one", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(str(docx_path))

    html_out = docx_to_html(str(docx_path))
    assert "<h1>Title</h1>" in html_out
    assert "<b>bold</b>" in html_out
    assert "<i>italic</i>" in html_out
    assert "<ul>" in html_out
    assert "<li>item one</li>" in html_out
    assert "<table" in html_out
    assert "<td>A</td>" in html_out
