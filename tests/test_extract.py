"""Tests for document extraction."""

import os

import pytest

from humantype.extract import ExtractionError, extract_paragraphs


def _make_minimal_pdf(path, lines):
    """Write a tiny valid one-page PDF with the given text lines."""
    content = b""
    for i, line in enumerate(lines):
        y = 720 - i * 20
        if line == "":
            line = " "  # keep a blank line so pdfplumber preserves the break
        safe = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"BT /F1 12 Tf 72 {y} Td ({safe}) Tj ET\n".encode("latin-1", "replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    with open(path, "wb") as fh:
        fh.write(out)


def test_extract_markdown(tmp_path):
    md = tmp_path / "sample.md"
    md.write_text("# Title\n\nSome **bold** text and [a link](http://x).\n\n- item\n",
                  encoding="utf-8")
    paragraphs = extract_paragraphs(str(md))
    assert paragraphs == ["Title", "Some bold text and a link.", "item"]


def test_extract_docx(tmp_path):
    from docx import Document
    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph with a quote \u2014 yes.")
    document.save(str(docx_path))

    paragraphs = extract_paragraphs(str(docx_path))
    assert paragraphs[0] == "First paragraph."
    assert paragraphs[1] == 'Second paragraph with a quote - yes.'


def test_extract_pdf(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    _make_minimal_pdf(str(pdf_path), ["Hello world", "", "Second paragraph"])
    paragraphs = extract_paragraphs(str(pdf_path))
    assert paragraphs[0] == "Hello world"
    assert paragraphs[1] == "Second paragraph"


def test_extract_missing_file():
    with pytest.raises(ExtractionError):
        extract_paragraphs("does_not_exist.pdf")


def test_extract_unsupported_extension(tmp_path):
    bad = tmp_path / "sample.xyz"
    bad.write_text("hi", encoding="utf-8")
    with pytest.raises(ExtractionError):
        extract_paragraphs(str(bad))


def test_extract_keeps_originals_docx(tmp_path):
    from docx import Document
    docx_path = tmp_path / "quote.docx"
    document = Document()
    document.add_paragraph("\u201cSmart quotes\u201d")
    document.save(str(docx_path))
    paragraphs = extract_paragraphs(str(docx_path), keep_originals=True)
    assert paragraphs[0] == "\u201cSmart quotes\u201d"
