"""Extract text from .pdf, .docx and .md files into a list of paragraphs."""

from __future__ import annotations

import os
import re

from .normalize import normalize_paragraph, strip_markdown, text_to_paragraphs


class ExtractionError(Exception):
    """Raised when a document cannot be extracted."""


def extract_paragraphs(path: str, keep_originals: bool = False) -> list[str]:
    """Extract paragraphs from a source document.

    Returns a list of paragraph strings with normalized whitespace.
    """
    if not os.path.exists(path):
        raise ExtractionError(f"file not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        paragraphs = _extract_pdf(path)
    elif ext == ".docx":
        paragraphs = _extract_docx(path)
    elif ext == ".doc":
        paragraphs = _extract_doc(path)
    elif ext in (".md", ".markdown", ".txt"):
        paragraphs = _extract_markdown_or_text(path)
    else:
        raise ExtractionError(
            f"unsupported file type {ext!r}; use .pdf, .doc, .docx, or .md"
        )
    paragraphs = [normalize_paragraph(p, keep_originals) for p in paragraphs]
    return [p for p in paragraphs if p]


def _extract_doc(path: str) -> list[str]:
    """Extract text from a legacy .doc by converting it with Word COM."""
    from . import word_com

    if not word_com.available():
        raise ExtractionError(
            "legacy .doc files need Word automation; install pywin32 "
            "(pip install pywin32) on Windows with Word installed"
        )
    converted = word_com.convert_to_docx(path)
    return _extract_docx(converted)


def _extract_pdf(path: str) -> list[str]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(
            "pdfplumber is required for PDF extraction; pip install pdfplumber"
        ) from exc

    paragraphs: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                # Prefer text lines with coordinates so we can detect paragraph
                # breaks from vertical gaps; fall back to plain text.
                try:
                    lines = page.extract_text_lines()
                except Exception:
                    lines = None
                if lines:
                    paragraphs.extend(_group_pdf_lines(lines))
                else:
                    text = page.extract_text(layout=False) or ""
                    page_paragraphs = re.split(r"\n\s*\n", text)
                    for block in page_paragraphs:
                        block = re.sub(r"\s*\n\s*", " ", block).strip()
                        if block:
                            paragraphs.append(block)
    except Exception as exc:
        raise ExtractionError(f"failed to read PDF {path}: {exc}") from exc
    if not paragraphs:
        raise ExtractionError(f"no text found in PDF: {path}")
    return paragraphs


def _group_pdf_lines(lines) -> list[str]:
    """Group pdfplumber text lines into paragraphs using vertical gaps."""
    def field(line, name, default=None):
        if isinstance(line, dict):
            return line.get(name, default)
        return getattr(line, name, default)

    heights = []
    for line in lines:
        top = field(line, "top")
        bottom = field(line, "bottom")
        if top is not None and bottom is not None:
            heights.append(bottom - top)
    median_height = float(sorted(heights)[len(heights) // 2]) if heights else 12.0
    gap_threshold = max(8.0, median_height * 1.2)

    paragraphs = []
    current = []
    previous = None
    for line in lines:
        text = (field(line, "text") or "").strip()
        if not text:
            continue
        top = field(line, "top")
        if previous is not None and top is not None:
            gap = top - previous
            if gap > gap_threshold:
                if current:
                    paragraphs.append(" ".join(current))
                    current = []
        current.append(text)
        previous = field(line, "bottom")
    if current:
        paragraphs.append(" ".join(current))
    return [p for p in paragraphs if p]


def _extract_docx(path: str) -> list[str]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(
            "python-docx is required for DOCX extraction; pip install python-docx"
        ) from exc

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ExtractionError(f"failed to read DOCX {path}: {exc}") from exc

    paragraphs = [p.text for p in document.paragraphs]
    # Include table cell text so nothing is silently dropped.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        paragraphs.append(p.text)
    paragraphs = [p for p in paragraphs if p.strip()]
    if not paragraphs:
        raise ExtractionError(f"no text found in DOCX: {path}")
    return paragraphs


def _extract_markdown_or_text(path: str) -> list[str]:
    with open(path, "r", encoding="utf-8-sig") as fh:
        text = fh.read()
    if os.path.splitext(path)[1].lower() == ".txt":
        return text_to_paragraphs(text)
    return text_to_paragraphs(strip_markdown(text))
