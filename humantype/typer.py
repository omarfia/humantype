"""Typing engine: plans a human-like keystroke stream and sends it.

The planner (TypingPlanner) generates a full event stream with timestamps and
keeps a private model of the document buffer, so corrections are guaranteed to
produce the exact source text when the uncorrected rate is 0.

The senders only translate events into real key presses; all timing decisions
happen in the planner.
"""

from __future__ import annotations

import csv
import datetime as _dt
import hashlib
import json
import math
import os
import random
import re
import sys
import time

from . import keymaps
from .errors import ErrorSpec, TypoEngine
from .timing import TimingModel

# --------------------------------------------------------------------------- #
# Event helpers
# --------------------------------------------------------------------------- #

EVENT_ORDER = {"keydown": 0, "keyup": 1, "marker": 2, "pause": 3}

CSV_FIELDS = ["timestamp", "elapsed", "key", "delay", "event_type", "duration", "detail"]


def split_token(token: str) -> tuple[str, str, str]:
    """Split a whitespace-free token into (prefix_punct, core, suffix_punct)."""
    match = re.match(r"^([^A-Za-z0-9]*)([A-Za-z0-9]*)([^A-Za-z0-9]*)$", token)
    if not match:
        return ("", token, "")
    return match.group(1), match.group(2), match.group(3)


# --------------------------------------------------------------------------- #
# Planner
# --------------------------------------------------------------------------- #

class TypingPlanner:
    """Generate the full keystroke event stream for a list of paragraphs."""

    def __init__(self, cfg, paragraphs: list[str], rng: random.Random | None = None):
        self.cfg = cfg
        self.paragraphs = paragraphs
        self.rng = rng or random.Random(cfg.typing.seed)
        self.timing = TimingModel(cfg, self.rng)
        self.engine = TypoEngine(cfg, self.rng)

        self.events: list[dict] = []
        self.buffer = ""
        self.cursor = 0
        self.now = 0.0
        self.last_keydown_t: float | None = None
        self.last_char: str | None = None
        self.paragraph_index = -1
        self.typed_words: list[str] = []
        self.last_save_at = 0.0
        self.save_interval = cfg.logging.save_interval_minutes * 60.0
        idle_rate = cfg.pauses.idle_rate_per_hour
        if idle_rate > 0:
            self.next_idle_at = self.rng.uniform(300.0, 1800.0)
        else:
            self.next_idle_at = float("inf")

        # Scripted mistakes (optional): a fixed list of typos instead of random
        # error generation.
        self._scripted = self._load_scripted()
        self._scripted_maps_built = False
        self._scripted_by_para: dict[tuple[int, int], ErrorSpec] = {}
        self._scripted_by_global: dict[int, ErrorSpec] = {}
        self._scripted_offsets: list[int] = []
        self._scripted_word_lists: list[list[str]] = []

    # ------------------------------------------------------------ scripted
    def _load_scripted(self):
        raw = getattr(self.cfg.errors, "scripted_mistakes", None)
        if raw is None:
            return None
        if isinstance(raw, str):
            import json
            with open(raw, "r", encoding="utf-8") as fh:
                data = json.load(fh)
        else:
            data = raw
        if not isinstance(data, list):
            raise ValueError("scripted_mistakes must be a JSON file containing a list, or an inline list")
        from .errors import error_spec_from_dict
        return [(item, error_spec_from_dict(item)) for item in data]

    def _ensure_scripted_maps(self):
        if self._scripted_maps_built or self._scripted is None:
            self._scripted_maps_built = True
            return
        word_lists = [[w for w in p.split(" ") if w] for p in self.paragraphs]
        self._scripted_word_lists = word_lists
        offsets = []
        count = 0
        for words in word_lists:
            offsets.append(count)
            count += len(words)
        self._scripted_offsets = offsets
        self._scripted_by_para = {}
        self._scripted_by_global = {}
        for item, spec in self._scripted:
            if "word_index" in item:
                self._scripted_by_global[int(item["word_index"])] = spec
            elif "paragraph" in item and "word" in item:
                key = (int(item["paragraph"]) - 1, int(item["word"]) - 1)
                self._scripted_by_para[key] = spec
            elif "word_text" in item:
                text = item["word_text"]
                for pi, words in enumerate(word_lists):
                    for wi, word in enumerate(words):
                        if word == text:
                            self._scripted_by_para[(pi, wi)] = spec
                            break
                    else:
                        continue
                    break
                else:
                    raise ValueError(f"scripted mistake word_text {text!r} not found in source")
            else:
                raise ValueError(
                    "each scripted mistake needs 'word_index', 'paragraph'+'word', or 'word_text'"
                )
        self._scripted_maps_built = True

    def _scripted_for(self, paragraph_index: int, word_index: int) -> ErrorSpec | None:
        if self._scripted is None:
            return None
        self._ensure_scripted_maps()
        key = (paragraph_index, word_index)
        if key in self._scripted_by_para:
            return self._scripted_by_para[key]
        global_index = self._scripted_offsets[paragraph_index] + word_index if self._scripted_offsets else word_index
        return self._scripted_by_global.get(global_index)

    # ------------------------------------------------------------------ plan
    def plan(self) -> list[dict]:
        sessions = max(1, int(self.cfg.sessions.count))
        groups = self._split_sessions(self.paragraphs, sessions)
        for group_index, group in enumerate(groups):
            if group_index > 0:
                self._pause("break", self._session_break_minutes(group_index) * 60.0)
            for paragraph in group:
                self._type_paragraph(paragraph)
        self._sort_events()
        return self.events

    @staticmethod
    def _split_sessions(paragraphs: list[str], sessions: int) -> list[list[str]]:
        if sessions <= 1 or not paragraphs:
            return [paragraphs]
        size = math.ceil(len(paragraphs) / sessions)
        return [paragraphs[i:i + size] for i in range(0, len(paragraphs), size)]

    def _session_break_minutes(self, group_index: int) -> float:
        schedule = list(self.cfg.sessions.schedule or [])
        if group_index - 1 < len(schedule):
            return float(schedule[group_index - 1])
        return 10.0

    # ------------------------------------------------------------- paragraphs
    def _type_paragraph(self, paragraph: str):
        self.paragraph_index += 1
        self.typed_words = []
        words = [w for w in paragraph.split(" ") if w != ""]
        if not words:
            return

        chunk_sizes = self.timing.chunk_sizes(len(words))
        word_in_burst: list[bool] = []
        for size in chunk_sizes:
            word_in_burst.extend([size >= 5] * size)
        chunk_ends = set()
        total = 0
        for size in chunk_sizes:
            total += size
            chunk_ends.add(total)

        pending_fix: ErrorSpec | None = None
        pending_fix_token = ""
        pending_fix_target = 0
        pending_fix_actual = 0

        i = 0
        while i < len(words):
            self._maybe_idle_gap()
            if (self._scripted is None and pending_fix is None
                    and self._maybe_lose_place(words, i, word_in_burst)):
                i += 2
                continue

            word = words[i]
            in_burst = word_in_burst[i]
            scripted = self._scripted_for(self.paragraph_index, i)

            if pending_fix is not None and pending_fix_actual >= pending_fix_target:
                self._late_correction(pending_fix_token, pending_fix_actual)
                pending_fix = None
                pending_fix_token = ""

            if pending_fix is not None:
                self._type_word_token(word, in_burst, allow_error=False)
                pending_fix_actual += 1
            else:
                err = self._type_word_token(word, in_burst, allow_error=True,
                                            force_error=scripted)
                if err is not None and err.late:
                    pending_fix = err
                    pending_fix_token = word
                    pending_fix_target = self.rng.randint(
                        1, max(1, self.cfg.errors.late_words_max)
                    )
                    pending_fix_actual = 0

            if i < len(words) - 1:
                next_first = words[i + 1][0] if words[i + 1] else None
                self._type_char(" ", event_type="normal", next_char=next_first, in_burst=False)
            self.typed_words.append(word)

            self._maybe_thinking_pause(word)

            if (i + 1) in chunk_ends and i + 1 < len(words):
                self._pause("chunk", self.timing.chunk_pause())
            i += 1

        if pending_fix is not None:
            self._late_correction(pending_fix_token, pending_fix_actual)
            pending_fix = None

        if self.paragraph_index < len(self.paragraphs) - 1:
            self._enter()
            self._pause("paragraph", self.timing.paragraph_pause())
        self._marker(f"para:{self.paragraph_index}")
        self._maybe_save()

    def _type_word_token(self, token: str, in_burst: bool, allow_error: bool,
                         force_error: ErrorSpec | None = None) -> ErrorSpec | None:
        prefix, core, suffix = split_token(token)
        err = None

        if prefix:
            nxt = core[0] if core else (suffix[0] if suffix else None)
            self._type_run(list(prefix), "normal", next_char=nxt, in_burst=in_burst)

        if core:
            if allow_error:
                if force_error is not None:
                    err = force_error
                elif self._scripted is not None:
                    err = None  # scripted mode: only the listed mistakes happen
                else:
                    err = self.engine.maybe_error(core, in_burst)
            suffix_next = suffix[0] if suffix else None
            if err is None:
                self._type_run(list(core), "normal", next_char=suffix_next, in_burst=in_burst)
            elif err.late or err.leave:
                self._type_core_wrong(core, err, in_burst, next_char=suffix_next)
            else:
                self._type_core_immediate(core, err, in_burst, next_char=suffix_next)

        if suffix:
            self._type_run(list(suffix), "normal", next_char=None, in_burst=in_burst)

        if err is not None and err.late:
            return err
        return None

    # ------------------------------------------------------------ core typing
    def _type_core_wrong(self, core: str, err: ErrorSpec, in_burst: bool, next_char=None):
        pairs = TypoEngine.wrong_chars(core, err)
        chars = [c for c, _ in pairs]
        for idx, (ch, is_typo) in enumerate(pairs):
            nxt = chars[idx + 1] if idx + 1 < len(chars) else next_char
            self._type_char(ch, event_type="typo" if is_typo else "normal",
                            next_char=nxt, in_burst=in_burst)

    def _type_core_immediate(self, core: str, err: ErrorSpec, in_burst: bool, next_char=None):
        p = err.pos
        k = err.notice_after
        t = err.type

        if t == "substitution":
            wrong_seg = [err.wrong_char]
            after_chars = list(core[p + 1:p + 1 + k])
            n_back = k + 1
            correct_seg = list(core[p:p + 1 + k])
            rest = core[p + 1 + k:]
        elif t == "transposition":
            wrong_seg = [core[p + 1], core[p]]
            after_chars = list(core[p + 2:p + 2 + k])
            n_back = k + 2
            correct_seg = list(core[p:p + 2 + k])
            rest = core[p + 2 + k:]
        elif t == "doubled":
            wrong_seg = [core[p], core[p]]
            after_chars = list(core[p + 1:p + 1 + k])
            n_back = k + 1
            correct_seg = list(core[p + 1:p + 1 + k])
            rest = core[p + 1 + k:]
        elif t == "dropped":
            wrong_seg = []
            after_chars = list(core[p + 1:p + 1 + k])
            n_back = len(after_chars)
            correct_seg = list(core[p:p + 1 + k])
            rest = core[p + 1 + k:]
        elif t == "shift_late":
            wrong_seg = [core[0], core[1].upper()]
            after_chars = list(core[2:2 + k])
            n_back = k + 2
            correct_seg = list(core[0:2 + k])
            rest = core[2 + k:]
        else:  # pragma: no cover - defensive
            self._type_run(list(core), "normal", next_char=next_char, in_burst=in_burst)
            return

        # Type up to the error, then the error itself.
        first_after = wrong_seg[0] if wrong_seg else (after_chars[0] if after_chars else None)
        self._type_run(list(core[:p]), "normal", next_char=first_after, in_burst=in_burst)
        wrong_next = after_chars[0] if after_chars else next_char
        for idx, ch in enumerate(wrong_seg):
            nxt = wrong_seg[idx + 1] if idx + 1 < len(wrong_seg) else wrong_next
            self._type_char(ch, "typo", next_char=nxt, in_burst=in_burst)
        self._type_run(after_chars, "normal", next_char=None, in_burst=in_burst)

        # Notice the error, backspace, pause, retype.
        for _ in range(n_back):
            self._backspace()
        self._pause("correction", self.timing.correction_pause())
        rest_first = rest[0] if rest else next_char
        self._type_run(correct_seg, "correction", next_char=rest_first, in_burst=in_burst)
        self._type_run(list(rest), "normal", next_char=next_char, in_burst=in_burst)

    # ------------------------------------------------------- corrections / nav
    def _late_correction(self, correct_token: str, words_typed: int):
        presses = words_typed + 1
        for _ in range(presses):
            self._combo(["Key.ctrl"], "Key.left", event_type="nav")
            self._move_cursor_word_left()
        self._combo(["Key.ctrl", "Key.shift"], "Key.right", event_type="nav")
        self._delete_word_at_cursor()
        self._pause("correction", self.timing.correction_pause())
        self._type_run(list(correct_token), "correction", next_char=None, in_burst=False)
        self._tap_key("Key.end", event_type="nav")
        self.cursor = len(self.buffer)

    def _maybe_lose_place(self, words: list[str], i: int, word_in_burst: list[bool]) -> bool:
        """Rarely lose your place.  Returns True when words i and i+1 are consumed."""
        if i == 0 or not self.typed_words:
            return False
        if self.rng.random() > self.cfg.errors.lose_place_probability:
            return False

        if self.rng.random() < 0.7 or i >= len(words) - 1:
            # Pause, look back, delete and retype the last word or two.
            m = self.rng.randint(1, min(2, len(self.typed_words)))
            self._pause("think", self.rng.uniform(2.0, 5.0))
            suffix = self._suffix_of_last_words(m)
            for _ in range(len(suffix)):
                self._backspace()
            self._pause("correction", self.timing.correction_pause())
            space_next = words[i][0] if words[i] else None
            self._type_run(list(suffix), "correction", next_char=space_next, in_burst=False)
            return False

        # Skip the current word, type the next one, then go back and insert it.
        skipped = words[i]
        next_word = words[i + 1]
        self._pause("think", self.rng.uniform(2.0, 5.0))
        self._type_word_token(next_word, word_in_burst[i + 1], allow_error=False)
        if i + 1 < len(words) - 1:
            self._type_char(" ", "normal", next_char=None, in_burst=False)
        self.typed_words.append(next_word)
        self._pause("think", self.rng.uniform(1.0, 2.5))
        self._combo(["Key.ctrl"], "Key.left", event_type="nav")
        self._move_cursor_word_left()
        self._type_run(list(skipped), "correction", next_char=" ", in_burst=False)
        self._type_char(" ", "correction", next_char=next_word[0] if next_word else None,
                        in_burst=False)
        self._tap_key("Key.end", event_type="nav")
        self.cursor = len(self.buffer)
        self.typed_words.insert(len(self.typed_words) - 1, skipped)
        return True

    # ---------------------------------------------------------------- buffers
    def _insert(self, ch: str):
        self.buffer = self.buffer[:self.cursor] + ch + self.buffer[self.cursor:]
        self.cursor += 1

    def _delete_before(self):
        if self.cursor > 0:
            self.buffer = self.buffer[:self.cursor - 1] + self.buffer[self.cursor:]
            self.cursor -= 1

    def _move_cursor_word_left(self):
        i = self.cursor
        while i > 0 and self.buffer[i - 1] in " \n":
            i -= 1
        while i > 0 and self.buffer[i - 1] not in " \n":
            i -= 1
        self.cursor = i

    def _delete_word_at_cursor(self):
        i = self.cursor
        while i < len(self.buffer) and self.buffer[i] not in " \n":
            i += 1
        self.buffer = self.buffer[:self.cursor] + self.buffer[i:]

    def _suffix_of_last_words(self, m: int) -> str:
        """Return the trailing text covering the last `m` typed words."""
        i = len(self.buffer)
        while i > 0 and self.buffer[i - 1] in " \n":
            i -= 1
        for step in range(m):
            while i > 0 and self.buffer[i - 1] not in " \n":
                i -= 1
            if step < m - 1:
                while i > 0 and self.buffer[i - 1] in " \n":
                    i -= 1
        return self.buffer[i:]

    # ------------------------------------------------------------------ emits
    def _emit(self, etype, key, t, event_type, delay, dur, detail=""):
        self.events.append({
            "t": round(t, 6),
            "type": etype,
            "key": key,
            "event_type": event_type,
            "delay": round(delay, 6),
            "duration": round(dur, 6),
            "detail": detail,
            "para": self.paragraph_index,
        })

    def _sort_events(self):
        self.events.sort(key=lambda e: (e["t"], EVENT_ORDER.get(e["type"], 9)))

    def _pause(self, label: str, duration: float):
        if duration <= 0:
            return
        self._emit("pause", "", self.now, "pause", 0.0, duration, label)
        self.now += duration
        if (label in ("idle", "break") and duration >= 30.0
                and self.save_interval > 0):
            self._combo(["Key.ctrl"], "s", event_type="save")
            self.last_save_at = self.now

    def _marker(self, detail: str):
        self._emit("marker", "", self.now, "para_end", 0.0, 0.0, detail)

    def _error_cluster(self) -> bool:
        return self.engine.words_since_error <= self.cfg.errors.error_cluster_window

    def _type_char(self, ch: str, event_type: str = "normal", next_char=None, in_burst: bool = False):
        if ch == "\n":
            self._enter()
            return
        base, needs_shift = keymaps.char_to_key(ch)
        t_down = self.now
        delay = 0.0 if self.last_keydown_t is None else max(0.0, t_down - self.last_keydown_t)
        hold = self.timing.hold_time()

        shift_up_t = None
        if needs_shift:
            shift_pre = self.rng.uniform(0.03, 0.08)
            self._emit("keydown", "Key.shift", t_down - shift_pre, "shift", 0.0, hold, "shift")
            shift_up_t = t_down + hold + self.rng.uniform(0.0, 0.05)

        self._emit("keydown", base, t_down, event_type, delay, hold, f"char:{ch}")
        self._emit("keyup", base, t_down + hold, "keyup", 0.0, hold, f"char:{ch}")
        if needs_shift:
            self._emit("keyup", "Key.shift", shift_up_t, "keyup", 0.0, hold, "shift")

        self._insert(ch)
        self.last_keydown_t = t_down
        self.last_char = ch
        self.now = t_down + self.timing.inter_key_delay(
            t_down, ch, next_char, in_burst, self._error_cluster()
        )

    def _type_run(self, chars: list[str], event_type: str, next_char=None, in_burst: bool = False):
        for idx, ch in enumerate(chars):
            nxt = chars[idx + 1] if idx + 1 < len(chars) else next_char
            self._type_char(ch, event_type=event_type, next_char=nxt, in_burst=in_burst)

    def _tap_key(self, key: str, event_type: str = "normal"):
        t_down = self.now
        delay = 0.0 if self.last_keydown_t is None else max(0.0, t_down - self.last_keydown_t)
        hold = self.timing.hold_time()
        self._emit("keydown", key, t_down, event_type, delay, hold, "")
        self._emit("keyup", key, t_down + hold, "keyup", 0.0, hold, "")
        self.last_keydown_t = t_down
        self.now = t_down + self.timing.tap_delay(key)

    def _backspace(self):
        self._tap_key("Key.backspace", event_type="backspace")
        self._delete_before()

    def _enter(self):
        self._tap_key("Key.enter", event_type="paragraph")
        self._insert("\n")

    def _combo(self, mods: list[str], key: str, event_type: str = "nav"):
        t0 = self.now
        delay = 0.0 if self.last_keydown_t is None else max(0.0, t0 - self.last_keydown_t)
        hold = self.timing.hold_time()
        for offset, mod in enumerate(mods):
            self._emit("keydown", mod, t0 + offset * 0.01, event_type, delay if offset == 0 else 0.0,
                       hold, "")
        t_key = t0 + 0.01 * len(mods)
        self._emit("keydown", key, t_key, event_type, delay if not mods else 0.0, hold, "")
        self._emit("keyup", key, t_key + hold, "keyup", 0.0, hold, "")
        for offset, mod in enumerate(reversed(mods)):
            self._emit("keyup", mod, t_key + hold + 0.01 + offset * 0.01, "keyup", 0.0, hold, "")
        self.last_keydown_t = t_key
        self.now = t0 + self.timing.uniform(0.15, 0.35)

    # ---------------------------------------------------------------- pauses
    def _maybe_thinking_pause(self, word: str):
        rate = self.cfg.typing.thinking_pause_rate
        if rate <= 0:
            return
        prob = 1.0 - math.exp(-rate * max(1, len(word)))
        if self.rng.random() < prob:
            self._pause("think", self.timing.thinking_pause())

    def _maybe_idle_gap(self):
        if self.now >= self.next_idle_at:
            self._pause("idle", self.timing.idle_pause())
            idle_rate = self.cfg.pauses.idle_rate_per_hour
            if idle_rate > 0:
                self.next_idle_at = self.now + self.rng.expovariate(idle_rate / 3600.0)
            else:
                self.next_idle_at = float("inf")

    def _maybe_save(self):
        if self.save_interval <= 0:
            return
        jitter = self.rng.uniform(0.8, 1.2)
        if self.now - self.last_save_at >= self.save_interval * jitter:
            self._combo(["Key.ctrl"], "s", event_type="save")
            self.last_save_at = self.now


# --------------------------------------------------------------------------- #
# Senders
# --------------------------------------------------------------------------- #

class PynputSender:
    def __init__(self):
        from pynput.keyboard import Controller, Key
        self.controller = Controller()
        self.Key = Key

    def send(self, ev: dict):
        if ev["type"] not in ("keydown", "keyup"):
            return
        key = self._to_key(ev["key"])
        if ev["type"] == "keydown":
            self.controller.press(key)
        else:
            self.controller.release(key)

    def _to_key(self, name: str):
        if len(name) == 1:
            return name
        attr = name.split(".", 1)[-1]
        return getattr(self.Key, attr, name)


class PyAutoGuiSender:
    def __init__(self):
        import pyautogui
        pyautogui.FAILSAFE = True  # corner failsafe is always on with pyautogui
        self.pyautogui = pyautogui

    def send(self, ev: dict):
        if ev["type"] not in ("keydown", "keyup"):
            return
        key = self._to_key(ev["key"])
        if ev["type"] == "keydown":
            self.pyautogui.keyDown(key)
        else:
            self.pyautogui.keyUp(key)

    @staticmethod
    def _to_key(name: str):
        if len(name) == 1:
            return name
        return name.split(".", 1)[-1]  # Key.space -> space, Key.ctrl -> ctrl


class CsvLog:
    def __init__(self, path: str):
        self.file = open(path, "w", newline="", encoding="utf-8")
        self.writer = csv.DictWriter(self.file, fieldnames=CSV_FIELDS)
        self.writer.writeheader()

    def write(self, ev: dict):
        row = {
            "timestamp": _dt.datetime.now().isoformat(timespec="milliseconds"),
            "elapsed": f"{ev['t']:.4f}",
            "key": ev["key"],
            "delay": f"{ev['delay']:.4f}" if ev["type"] == "keydown" else "",
            "event_type": ev["event_type"] if ev["type"] != "keyup" else "keyup",
            "duration": f"{ev['duration']:.4f}" if ev["duration"] else "",
            "detail": ev.get("detail", ""),
        }
        self.writer.writerow(row)
        self.file.flush()

    def close(self):
        self.file.close()


# --------------------------------------------------------------------------- #
# Run control (hotkeys, progress)
# --------------------------------------------------------------------------- #

class RunController:
    def __init__(self):
        self.paused = False
        self.aborted = False
        self._pause_notified = False


def start_hotkey_listener(cfg, controller: RunController):
    try:
        from pynput import keyboard
    except ImportError:
        return None
    pause_key = getattr(keyboard.Key, cfg.hotkeys.pause.lower(), keyboard.Key.f8)
    abort_key = getattr(keyboard.Key, cfg.hotkeys.abort.lower(), keyboard.Key.f9)

    def on_press(key):
        if key == pause_key:
            controller.paused = not controller.paused
            controller._pause_notified = False
        elif key == abort_key:
            controller.aborted = True

    listener = keyboard.Listener(on_press=on_press)
    listener.daemon = True
    listener.start()
    return listener


def source_sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def load_progress(cfg, source_path: str, sha: str) -> int:
    path = cfg.logging.progress_file
    if not os.path.exists(path):
        return 0
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, json.JSONDecodeError):
        return 0
    if data.get("source_sha256") != sha or data.get("source_path") != os.path.abspath(source_path):
        print(f"warning: progress file {path} does not match this source; starting from the top")
        return 0
    return int(data.get("last_paragraph_index", -1)) + 1


def save_progress(cfg, source_path: str, sha: str, paragraph_index: int, total: int):
    path = cfg.logging.progress_file
    data = {
        "source_path": os.path.abspath(source_path),
        "source_sha256": sha,
        "last_paragraph_index": paragraph_index,
        "total_paragraphs": total,
        "updated_at": _dt.datetime.now().isoformat(timespec="seconds"),
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2)


def send_ctrl_s(sender):
    """Send a plain Ctrl+S (used after typing to save the document)."""
    for ev in (
        {"type": "keydown", "key": "Key.ctrl"},
        {"type": "keydown", "key": "s"},
        {"type": "keyup", "key": "s"},
        {"type": "keyup", "key": "Key.ctrl"},
    ):
        sender.send(ev)
        time.sleep(0.05)


# --------------------------------------------------------------------------- #
# Dry run printing
# --------------------------------------------------------------------------- #

def print_event(ev: dict):
    t = ev["t"]
    if ev["type"] == "pause":
        print(f"{t:9.3f}s  pause      {ev['detail']:<12} {ev['duration']:6.2f}s")
    elif ev["type"] == "marker":
        print(f"{t:9.3f}s  --- paragraph {ev['detail'].split(':', 1)[-1]} complete ---")
    elif ev["type"] == "keydown":
        hold = ev["duration"] * 1000
        print(
            f"{t:9.3f}s  keydown  {ev['key']:<16} "
            f"delay={ev['delay']:6.3f}s  hold={hold:5.1f}ms  {ev['event_type']:<10} {ev['detail']}"
        )
    else:
        pass  # keyup events are not printed


# --------------------------------------------------------------------------- #
# Main run loop
# --------------------------------------------------------------------------- #

def _say(status_cb, message: str):
    if status_cb is not None:
        status_cb(message)
    else:
        print(message)


def run_typing(cfg, paragraphs: list[str], source_path: str,
               resume: bool = False, verify_docx: str | None = None,
               controller: RunController | None = None,
               progress_cb=None, status_cb=None,
               send_final_save: bool = True,
               progress_offset: int = 0):
    """Plan and execute a typing run.

    `progress_cb(paragraph_index, total)` fires after each completed paragraph
    (both values are full-document numbers when `progress_offset` is used).
    `status_cb(message)` receives human-readable status updates.
    `send_final_save` controls the Ctrl+S that is pressed before verification;
    set it to False when the target document may be unsaved (Ctrl+S would open
    the Save As dialog).
    Returns {"aborted": bool, "log_path": str|None, "buffer_matches": bool}.
    """
    dry_run = getattr(cfg.typing, "dry_run", False)
    rng = random.Random(cfg.typing.seed)
    planner = TypingPlanner(cfg, paragraphs, rng)
    events = planner.plan()

    total = len(paragraphs)
    full_total = total + progress_offset
    sha = source_sha256(source_path) if os.path.exists(source_path) else ""
    start_para = load_progress(cfg, source_path, sha) if resume else 0
    start_para = max(0, start_para - progress_offset)
    if resume and start_para > 0:
        _say(status_cb, f"resuming from paragraph {start_para + progress_offset + 1} of {full_total}")

    active = [e for e in events if e.get("para", 0) >= start_para]
    if not active:
        _say(status_cb, "nothing to type: all selected paragraphs already completed")
        return {"aborted": False, "log_path": None, "buffer_matches": True}

    t0 = active[0]["t"]

    if dry_run:
        limit = cfg.logging.dry_run_max_events
        shown = 0
        for ev in active:
            if limit and shown >= limit:
                _say(status_cb, f"... ({len(active) - shown} more events)")
                break
            print_event(ev)
            shown += 1
        _say(status_cb, f"\nplanned {len(active)} events, ~{active[-1]['t'] - t0:.1f}s of typing")
        _say(status_cb, f"final buffer matches source: {planner.buffer == chr(10).join(paragraphs)}")
        return {"aborted": False, "log_path": None,
                "buffer_matches": planner.buffer == chr(10).join(paragraphs)}

    # Choose a sender.
    backend = cfg.typing.backend.lower()
    if backend == "pyautogui":
        sender = PyAutoGuiSender()
    else:
        sender = PynputSender()

    own_controller = controller is None
    if controller is None:
        controller = RunController()
    listener = start_hotkey_listener(cfg, controller) if own_controller else None

    log_path = cfg.logging.output
    if log_path is None:
        stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        log_path = f"humantype_log_{stamp}.csv"
    log = CsvLog(log_path)
    _say(status_cb, f"logging events to {log_path}")

    countdown = cfg.typing.countdown
    _say(status_cb, f"Switch to your Word window now.  Starting in {countdown:g} seconds...")
    _say(status_cb, f"  {cfg.hotkeys.pause} = pause/resume, {cfg.hotkeys.abort} = abort")
    end_at = time.time() + countdown
    while time.time() < end_at:
        if controller.aborted:
            break
        remaining = end_at - time.time()
        if remaining > 0.05:
            _say(status_cb, f"starting in {remaining:4.1f}s")
            time.sleep(min(0.1, remaining))
    if controller.aborted:
        _say(status_cb, "aborted during countdown")
        if listener:
            listener.stop()
        log.close()
        return {"aborted": True, "log_path": log_path, "buffer_matches": False}
    _say(status_cb, "typing...")

    start_wall = time.time()
    last_completed = start_para - 1
    aborted = False
    try:
        for ev in active:
            if controller.aborted:
                aborted = True
                break
            while controller.paused and not controller.aborted:
                if not controller._pause_notified:
                    _say(status_cb, "paused (F8 to resume, F9 to abort)")
                    controller._pause_notified = True
                time.sleep(0.2)
            if controller.aborted:
                aborted = True
                break

            target = start_wall + (ev["t"] - t0)
            wait = target - time.time()
            if wait > 0:
                if ev["type"] == "pause" and ev["duration"] > 1.0:
                    _say(status_cb, f"... {ev['detail']} pause for {ev['duration']:.0f}s")
                time.sleep(wait)

            if ev["type"] == "marker":
                if ev.get("para", 0) > last_completed:
                    last_completed = ev["para"]
                    save_progress(cfg, source_path, sha,
                                  last_completed + progress_offset, full_total)
                    if progress_cb is not None:
                        progress_cb(last_completed + progress_offset, full_total)
                continue
            sender.send(ev)
            log.write(ev)
    finally:
        if listener:
            listener.stop()
        log.close()

    if aborted:
        _say(status_cb,
             f"aborted.  Progress saved: completed through paragraph "
             f"{last_completed + progress_offset + 1} of {full_total}")
        _say(status_cb, f"resume later with: humantype type {source_path} --resume")
    else:
        save_progress(cfg, source_path, sha, total - 1 + progress_offset, full_total)
        _say(status_cb, f"finished typing {total} paragraphs")
        if progress_cb is not None:
            progress_cb(total - 1 + progress_offset, full_total)

    if not aborted and verify_docx:
        if send_final_save:
            _say(status_cb, "saving document (Ctrl+S)...")
            send_ctrl_s(sender)
            time.sleep(2.0)
        else:
            _say(status_cb, "skipping Ctrl+S (document may be unsaved)")
        compare_with_source(verify_docx, paragraphs)

    return {"aborted": aborted, "log_path": log_path,
            "buffer_matches": planner.buffer == chr(10).join(paragraphs)}


def run_typing_com(cfg, paragraphs: list[str], source_path: str, target_doc,
                   resume: bool = False, controller: RunController | None = None,
                   progress_cb=None, status_cb=None, progress_offset: int = 0):
    """Background typing via Word COM.

    Reuses the exact same human-like event plan as `run_typing` (same delays,
    pauses, typos and corrections) but inserts the text into `target_doc`
    through Word automation instead of sending keystrokes.  The Word window
    can stay in the background while the user works in other apps.
    """
    from . import word_com

    dry_run = getattr(cfg.typing, "dry_run", False)
    rng = random.Random(cfg.typing.seed)
    planner = TypingPlanner(cfg, paragraphs, rng)
    events = planner.plan()

    total = len(paragraphs)
    full_total = total + progress_offset
    sha = source_sha256(source_path) if os.path.exists(source_path) else ""
    start_para = load_progress(cfg, source_path, sha) if resume else 0
    start_para = max(0, start_para - progress_offset)

    active = [e for e in events if e.get("para", 0) >= start_para]
    if not active:
        _say(status_cb, "nothing to type: all selected paragraphs already completed")
        return {"aborted": False, "log_path": None, "buffer_matches": True}

    t0 = active[0]["t"]

    if dry_run:
        limit = cfg.logging.dry_run_max_events
        shown = 0
        for ev in active:
            if limit and shown >= limit:
                _say(status_cb, f"... ({len(active) - shown} more events)")
                break
            print_event(ev)
            shown += 1
        _say(status_cb, f"\nplanned {len(active)} events, ~{active[-1]['t'] - t0:.1f}s of typing")
        _say(status_cb, f"final buffer matches source: {planner.buffer == chr(10).join(paragraphs)}")
        return {"aborted": False, "log_path": None,
                "buffer_matches": planner.buffer == chr(10).join(paragraphs)}

    log_path = cfg.logging.output
    if log_path is None:
        stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        log_path = f"humantype_log_{stamp}.csv"
    log = CsvLog(log_path)
    _say(status_cb, f"background typing via Word COM; logging events to {log_path}")

    countdown = min(cfg.typing.countdown, 3.0)
    _say(status_cb,
         f"starting background typing in {countdown:g}s - you can keep using the computer")
    end_at = time.time() + countdown
    while time.time() < end_at:
        if controller.aborted:
            break
        time.sleep(0.1)
    if controller.aborted:
        _say(status_cb, "aborted during countdown")
        log.close()
        return {"aborted": True, "log_path": log_path, "buffer_matches": False}

    word = target_doc.Application
    try:
        word.Selection.EndKey(Unit=6)  # wdStory: move to end of document
    except Exception as exc:
        log.close()
        raise word_com.WordComError(f"could not position Word cursor: {exc}") from exc

    start_wall = time.time()
    last_completed = start_para - 1
    aborted = False
    ctrl_held = False
    shift_held = False
    try:
        for ev in active:
            if controller.aborted:
                aborted = True
                break
            while controller.paused and not controller.aborted:
                if not controller._pause_notified:
                    _say(status_cb, "paused (use the Pause/Resume button)")
                    controller._pause_notified = True
                time.sleep(0.2)
            if controller.aborted:
                aborted = True
                break

            target = start_wall + (ev["t"] - t0)
            wait = target - time.time()
            if wait > 0:
                time.sleep(wait)

            if ev["type"] == "marker":
                if ev.get("para", 0) > last_completed:
                    last_completed = ev["para"]
                    save_progress(cfg, source_path, sha,
                                  last_completed + progress_offset, full_total)
                    if progress_cb is not None:
                        progress_cb(last_completed + progress_offset, full_total)
                continue
            if ev["type"] == "pause":
                log.write(ev)
                continue

            if ev["type"] == "keydown":
                key = ev["key"]
                if key == "Key.ctrl":
                    ctrl_held = True
                elif key == "Key.shift":
                    shift_held = True
                elif key == "Key.left" and ctrl_held:
                    word.Selection.MoveLeft(Unit=2, Count=1)  # wdWord
                elif key == "Key.right" and ctrl_held and shift_held:
                    word.Selection.MoveRight(Unit=2, Count=1, Extend=1)  # wdWord
                elif key == "Key.end":
                    word.Selection.EndKey(Unit=6)  # wdStory
                elif key == "Key.backspace":
                    word.Selection.TypeBackspace()
                elif key == "Key.enter":
                    word.Selection.TypeParagraph()
                elif key == "Key.space":
                    word.Selection.TypeText(" ")
                elif len(key) == 1 and ctrl_held and key.lower() == "s":
                    try:
                        if target_doc.Path:
                            target_doc.Save()
                    except Exception:
                        pass
                elif len(key) == 1:
                    ch = ev.get("detail", "")
                    if ch.startswith("char:"):
                        ch = ch[5:]
                    else:
                        ch = key
                    word.Selection.TypeText(ch)
            elif ev["type"] == "keyup":
                if ev["key"] == "Key.ctrl":
                    ctrl_held = False
                elif ev["key"] == "Key.shift":
                    shift_held = False
            log.write(ev)
    finally:
        log.close()

    if aborted:
        _say(status_cb,
             f"aborted.  Progress saved: completed through paragraph "
             f"{last_completed + progress_offset + 1} of {full_total}")
    else:
        save_progress(cfg, source_path, sha, total - 1 + progress_offset, full_total)
        _say(status_cb, f"finished background typing {total} paragraphs")
        if progress_cb is not None:
            progress_cb(total - 1 + progress_offset, full_total)

    return {"aborted": aborted, "log_path": log_path,
            "buffer_matches": planner.buffer == chr(10).join(paragraphs)}


def compare_with_source(docx_path: str, source_paragraphs: list[str]):
    """Read a saved .docx and report differences against the source text."""
    try:
        from docx import Document
        document = Document(docx_path)
        actual = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    except Exception as exc:
        print(f"could not read {docx_path} for verification: {exc}")
        return

    expected = [p.strip() for p in source_paragraphs if p.strip()]
    import difflib
    diff = list(difflib.unified_diff(expected, actual, fromfile="source", tofile="word", lineterm=""))
    if not diff:
        print("verification: Word document matches the source text exactly")
    else:
        print("verification: differences found (source vs Word document):")
        for line in diff[:40]:
            print("  " + line)
        if len(diff) > 40:
            print(f"  ... {len(diff) - 40} more diff lines")
